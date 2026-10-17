"""Export writers (spec section 11): CSV (streamed, UTF-8 BOM), XLSX, PNG, PDF.

Security: findings contain attacker-influenced text (titles, paths). Cells that a spreadsheet
would evaluate as formulas are neutralised (OWASP "CSV injection").
"""

from __future__ import annotations

import csv
import io
import logging
from collections.abc import Iterable, Iterator, Mapping, Sequence
from datetime import UTC, date, datetime
from typing import Any

from openpyxl import Workbook

log = logging.getLogger(__name__)

_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r", "\n")
EXPORT_COLUMNS = [
    "project", "branch", "engine", "severity", "state", "status", "title", "cwe", "cve", "cvss", "package", "package_version",
    "fixed_version", "image", "file_path", "line", "language", "category", "first_found_at", "scan_id", "finding_id",
]  # fmt: skip


def sanitize_cell(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.astimezone(UTC).strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str) and value.startswith(_FORMULA_PREFIXES):
        return "'" + value
    return value


def csv_stream(rows: Iterable[Mapping[str, Any]], columns: Sequence[str] = EXPORT_COLUMNS) -> Iterator[bytes]:
    """Yield CSV chunks; the first chunk starts with a BOM so Excel detects UTF-8."""
    buf = io.StringIO()
    writer = csv.writer(buf, lineterminator="\r\n")
    writer.writerow(columns)
    yield b"\xef\xbb\xbf" + buf.getvalue().encode()
    buf.seek(0)
    buf.truncate()
    for i, row in enumerate(rows, 1):
        writer.writerow([sanitize_cell(row.get(c)) for c in columns])
        if i % 500 == 0:
            yield buf.getvalue().encode()
            buf.seek(0)
            buf.truncate()
    if buf.tell():
        yield buf.getvalue().encode()


def xlsx_bytes(
    rows: Iterable[Mapping[str, Any]],
    summary: Mapping[str, Any],
    columns: Sequence[str] = EXPORT_COLUMNS,
    max_rows: int = 1_000_000,
) -> bytes:
    """Summary sheet (filters used, generated at) + Data sheet. Write-only mode keeps memory flat."""
    wb = Workbook(write_only=True)
    ws_sum = wb.create_sheet("Summary")
    ws_sum.append(["Generated at (UTC)", datetime.now(UTC).strftime("%Y-%m-%d %H:%M:%S")])
    for k, v in summary.items():
        ws_sum.append([str(k), sanitize_cell(", ".join(map(str, v)) if isinstance(v, (list, tuple)) else v)])
    ws = wb.create_sheet("Data")
    ws.append(list(columns))
    for i, row in enumerate(rows):
        if i >= min(max_rows, 1_048_575):  # Excel sheet limit
            break
        ws.append([sanitize_cell(row.get(c)) for c in columns])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def figure_to_png(fig: Any, width: int = 900, height: int = 500) -> bytes | None:
    """Server-side PNG via kaleido (optional dependency). The browser modebar covers the interactive case."""
    try:
        return fig.to_image(format="png", width=width, height=height, scale=2)  # type: ignore[no-any-return]
    except Exception as exc:  # kaleido missing or no Chrome available
        log.warning("png export unavailable", extra={"error": type(exc).__name__})
        return None


def _deny_fetch(url: str, *_: Any, **__: Any) -> Any:
    """The report is self-contained (inline CSS, data: images); never let the renderer fetch anything."""
    raise ValueError(f"external resource blocked: {url[:60]}")


def html_to_pdf(html: str) -> bytes | None:
    """HTML -> PDF with WeasyPrint (preferred) or Playwright; ``None`` when neither is installed."""
    try:
        from weasyprint import HTML

        return HTML(string=html, url_fetcher=_deny_fetch).write_pdf()  # type: ignore[no-any-return]
    except Exception as exc:
        log.debug("weasyprint unavailable", extra={"error": type(exc).__name__})
    try:
        from playwright.sync_api import sync_playwright

        with sync_playwright() as p:
            browser = p.chromium.launch()
            page = browser.new_context(java_script_enabled=False).new_page()
            page.set_content(html)
            pdf = page.pdf(format="A4", print_background=True)
            browser.close()
            return pdf
    except Exception as exc:
        log.warning("pdf export unavailable", extra={"error": type(exc).__name__})
        return None
